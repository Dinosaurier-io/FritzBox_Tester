"""Export eines Testlaufs als einzelne Excel-Arbeitsmappe.

Warum zusaetzlich zu CSV und JSON? Die beiden bestehenden Formate richten sich
an Werkzeuge: CSV an Auswertungsprogramme, JSON an Skripte. Wer den Lauf
dagegen einfach ansehen oder weitergeben will, bekam bisher acht Dateien mit
Namen wie ``run0003_measurements.csv`` - und musste sie einzeln oeffnen, um zu
sehen, was gemessen wurde.

Diese Arbeitsmappe ist eine Datei, benannt nach dem Testlauf, mit einer
lesbaren Uebersicht auf dem ersten Blatt und den Rohdaten dahinter. Die
Kennzahlen der Uebersicht stammen aus derselben Auswertung wie der
HTML-Bericht - zwei Wege zu denselben Zahlen waeren zwei Wege, sie
unterschiedlich falsch zu berechnen.

Das Blatt mit den Rohmesswerten ist begrenzt: Ein Lauf ueber drei Tage
erreicht mehrere hunderttausend Zeilen. Wird die Grenze ueberschritten, steht
das ausdruecklich in der Datei - stillschweigend abgeschnittene Messdaten
waeren in einem Messwerkzeug nicht vertretbar.
"""

from __future__ import annotations

import logging
import re
import sqlite3
from collections.abc import Iterable, Iterator, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fbtest.report.generator import analyze_run
from fbtest.storage.database import Database

log = logging.getLogger(__name__)

#: Hoechstzahl an Rohmesswerten in der Arbeitsmappe.
#:
#: Excel selbst verkraftet gut eine Million Zeilen, oeffnet eine solche Datei
#: aber traege. Wer mehr braucht, ist mit dem CSV-Export besser bedient - der
#: Hinweis darauf steht in der Datei.
MAX_RAW_ROWS = 200_000

#: Zeichen, die Windows in Dateinamen nicht zulaesst.
_FORBIDDEN = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

#: Reservierte Geraetenamen unter Windows - als Dateiname nicht verwendbar.
_RESERVED = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=13)


def safe_filename(name: str, fallback: str = "Testlauf") -> str:
    """Macht einen Testlaufnamen als Dateinamen verwendbar.

    Der Name kommt aus einem Eingabefeld und darf alles enthalten - auch
    Schraegstriche, Doppelpunkte oder nichts als Leerzeichen.

    Args:
        name: Vom Benutzer vergebener Name.
        fallback: Ersatz, wenn nichts Brauchbares uebrig bleibt.

    Returns:
        Ein Dateiname ohne Endung, der unter Windows und Linux zulaessig ist.
    """
    cleaned = _FORBIDDEN.sub("", name).strip().strip(".")
    cleaned = re.sub(r"\s+", " ", cleaned)[:80].strip()
    if not cleaned or cleaned.upper() in _RESERVED:
        return fallback
    return cleaned


def _unique_path(directory: Path, stem: str, suffix: str, run_id: int) -> Path:
    """Findet einen freien Dateinamen.

    Zwei Testlaeufe duerfen denselben Namen tragen - ein Export darf den
    anderen deshalb nicht ueberschreiben. Erst bei Gleichstand kommt die
    Laufnummer dazu, sonst hiesse jede Datei ``Test_0003``.
    """
    candidate = directory / f"{stem}{suffix}"
    if not candidate.exists():
        return candidate
    numbered = directory / f"{stem} (Lauf {run_id}){suffix}"
    if not numbered.exists():
        return numbered
    for counter in range(2, 100):
        alternative = directory / f"{stem} (Lauf {run_id}-{counter}){suffix}"
        if not alternative.exists():
            return alternative
    raise OSError(f"Kein freier Dateiname fuer '{stem}' in {directory}.")


def _local_time(timestamp: float | None) -> str:
    """Formatiert einen Zeitstempel als lokale Zeit."""
    if timestamp is None:
        return ""
    return datetime.fromtimestamp(timestamp, tz=UTC).astimezone().strftime("%d.%m.%Y %H:%M:%S")


def _write_header(sheet: Worksheet, columns: Sequence[str]) -> None:
    """Schreibt eine hervorgehobene Kopfzeile und friert sie ein."""
    sheet.append(list(columns))
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    # Als Zellbezug, nicht ueber sheet.cell(): Der Zugriff auf eine Zelle legt
    # sie an, und das Blatt haette danach eine leere Zeile unter der Kopfzeile.
    sheet.freeze_panes = f"A{sheet.max_row + 1}"


def _autosize(sheet: Worksheet, widths: Sequence[int]) -> None:
    """Setzt feste Spaltenbreiten.

    Bewusst keine Berechnung aus dem Inhalt: Dafuer muesste jede Zelle noch
    einmal gelesen werden, und bei zweihunderttausend Zeilen kostet das mehr
    Zeit als der gesamte uebrige Export.
    """
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _cell_value(value: Any) -> Any:
    """Bereitet einen Datenbankwert fuer eine Zelle auf."""
    if value is None:
        return ""
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def _sheet_from_rows(
    workbook: Workbook,
    title: str,
    rows: Iterable[sqlite3.Row],
    *,
    limit: int | None = None,
) -> int:
    """Legt ein Blatt aus Datenbankzeilen an.

    Args:
        workbook: Zielmappe.
        title: Blattname.
        rows: Zeilen der Tabelle.
        limit: Hoechstzahl geschriebener Zeilen; ``None`` fuer alle.

    Returns:
        Zahl der geschriebenen Datenzeilen.
    """
    sheet = workbook.create_sheet(title)
    written = 0
    iterator: Iterator[sqlite3.Row] = iter(rows)

    first = next(iterator, None)
    if first is None:
        sheet.append(["Keine Daten in diesem Testlauf."])
        return 0

    columns = list(first.keys())
    _write_header(sheet, columns)

    # Zeitstempel zusaetzlich als lesbare Ortszeit - die Unix-Sekunden bleiben
    # erhalten, weil sich nur mit ihnen weiterrechnen laesst.
    time_column = "ts" if "ts" in columns else ("started_at" if "started_at" in columns else None)
    if time_column is not None:
        sheet.cell(row=1, column=len(columns) + 1, value="zeit_lokal")
        header = sheet.cell(row=1, column=len(columns) + 1)
        header.fill = _HEADER_FILL
        header.font = _HEADER_FONT

    for row in (first, *iterator):
        if limit is not None and written >= limit:
            sheet.append(["... hier abgeschnitten. Vollstaendige Daten ueber den CSV-Export."])
            break
        values = [_cell_value(row[column]) for column in columns]
        if time_column is not None:
            values.append(_local_time(row[time_column]))
        sheet.append(values)
        written += 1

    _autosize(sheet, [14] * (len(columns) + 1))
    return written


def export_workbook(db: Database, run_id: int, target_dir: Path) -> Path:
    """Exportiert einen Testlauf als eine Excel-Arbeitsmappe.

    Args:
        db: Geoeffnete Datenbank.
        run_id: ID des Testlaufs.
        target_dir: Zielordner; wird bei Bedarf angelegt.

    Returns:
        Pfad der geschriebenen Datei.

    Raises:
        ValueError: Wenn der Testlauf nicht existiert.
    """
    run = db.get_test_run(run_id)
    if run is None:
        raise ValueError(f"Testlauf #{run_id} existiert nicht.")

    analysis = analyze_run(db, run_id)
    target_dir.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    # Die automatisch angelegte erste Tabelle wird durch die Uebersicht ersetzt.
    workbook.remove(workbook.active)

    _build_overview(workbook, db, run_id, analysis)
    _build_series(workbook, db, run_id)

    _sheet_from_rows(workbook, "Ereignisse", db.iter_table("events", run_id))
    _sheet_from_rows(workbook, "Ausfaelle", db.iter_table("outages", run_id))
    _sheet_from_rows(workbook, "Bandbreite", db.iter_table("speedtests", run_id))
    _sheet_from_rows(workbook, "Router-Status", db.iter_table("router_status", run_id))
    _sheet_from_rows(workbook, "WLAN-Status", db.iter_table("wlan_status", run_id))
    _sheet_from_rows(
        workbook, "Messwerte", db.iter_table("measurements", run_id), limit=MAX_RAW_ROWS
    )

    path = _unique_path(target_dir, safe_filename(run.name), ".xlsx", run_id)
    workbook.save(path)
    log.info("Arbeitsmappe geschrieben: %s", path)
    return path


def _build_overview(workbook: Workbook, db: Database, run_id: int, analysis: Any) -> None:
    """Baut das Uebersichtsblatt aus den Kennzahlen der Auswertung."""
    sheet = workbook.create_sheet("Uebersicht")
    run = analysis.run

    sheet.append([f"Testlauf: {run.name}"])
    sheet["A1"].font = _TITLE_FONT
    sheet.append([])

    for label, value in (
        ("Laufnummer", run.id),
        ("Gestartet", _local_time(run.started_at)),
        ("Beendet", _local_time(run.ended_at)),
        ("Laufzeit", analysis.duration_text),
        ("Router", run.router_model or "unbekannt"),
        ("Firmware", run.firmware_version or "unbekannt"),
        ("Notizen", run.notes or ""),
    ):
        sheet.append([label, value])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

    sheet.append([])
    sheet.append(["Kennzahl", "Wert", "Einheit", "Erlaeuterung"])
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for metric in analysis.summary:
        sheet.append([metric.label, metric.value, metric.unit, metric.hint])

    sheet.append([])
    sheet.append(["Datenumfang", "Zeilen"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for table, label in (
        ("measurements", "Messwerte"),
        ("events", "Ereignisse"),
        ("outages", "Ausfaelle"),
        ("speedtests", "Bandbreitenmessungen"),
        ("router_status", "Router-Abfragen"),
        ("wlan_status", "WLAN-Abfragen"),
    ):
        sheet.append([label, db.count_rows(table, run_id)])

    _autosize(sheet, [28, 22, 12, 70])
    sheet.column_dimensions["D"].width = 70
    for row in sheet.iter_rows(min_col=4, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _build_series(workbook: Workbook, db: Database, run_id: int) -> None:
    """Baut das Blatt mit den Kennzahlen je Messreihe.

    Median und 95%-Perzentil stehen bewusst neben dem Mittelwert: Bei Latenzen
    sagt der Mittelwert wenig aus, weil ihn einzelne Ausreisser verschieben.
    Erst der Abstand zwischen Median und Perzentil zeigt, wie gleichmaessig
    eine Verbindung war.
    """
    sheet = workbook.create_sheet("Messreihen")
    _write_header(
        sheet,
        ["Modul", "Messreihe", "Einheit", "Anzahl",
         "Minimum", "Mittel", "Median", "p95", "Maximum"],
    )

    for row in db.list_metric_series(run_id):
        module, metric, unit = row["module"], row["metric"], row["unit"]
        stats = db.metric_stats(run_id, module, metric)
        if stats is None:
            continue
        sheet.append([
            module,
            metric,
            unit,
            int(stats["count"]),
            round(stats["min"], 3),
            round(stats["avg"], 3),
            _rounded(db.percentile(run_id, module, metric, 50)),
            _rounded(db.percentile(run_id, module, metric, 95)),
            round(stats["max"], 3),
        ])

    if sheet.max_row == 1:
        sheet.append(["Keine Messwerte in diesem Testlauf."])
    _autosize(sheet, [14, 22, 10, 10, 12, 12, 12, 12, 12])


def _rounded(value: float | None) -> float | str:
    """Rundet einen Perzentilwert; leere Zelle, wenn es keinen gibt."""
    return "" if value is None else round(value, 3)
