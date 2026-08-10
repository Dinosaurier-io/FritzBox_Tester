"""Tests des Excel-Exports: eine Datei je Testlauf, benannt nach dem Lauf."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from fbtest.core.models import Event, EventType, Measurement, SpeedtestResult
from fbtest.storage import workbook as workbook_module
from fbtest.storage.database import Database
from fbtest.storage.workbook import export_workbook, safe_filename


def _filled_run(db: Database, name: str = "Firmware 8.02", count: int = 5) -> int:
    """Legt einen abgeschlossenen Testlauf mit Daten an."""
    run = db.create_test_run(name, "{}", "8.02", "FRITZ!Box 7590")
    db.insert_measurements(
        run.id,
        [Measurement("ping", "rtt_avg", 10.0 + i, "ms", timestamp=1000.0 + i) for i in range(count)],
    )
    db.insert_events(run.id, [Event(EventType.RUN_START, "Start")])
    db.insert_speedtests(run.id, [SpeedtestResult(down_mbps=95.0, up_mbps=40.0)])
    db.finish_test_run(run.id)
    return run.id


class TestSafeFilename:
    """Ein Testlaufname darf alles enthalten - ein Dateiname nicht."""

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            ("Firmware 8.02", "Firmware 8.02"),
            ("vor/nach Umbau", "vornach Umbau"),
            ("Test: WLAN?", "Test WLAN"),
            ("  mehrere   Leerzeichen ", "mehrere Leerzeichen"),
        ],
    )
    def test_cleans_the_name(self, value: str, expected: str) -> None:
        assert safe_filename(value) == expected

    @pytest.mark.parametrize("value", ["", "   ", "...", "///", "CON", "lpt1"])
    def test_falls_back_when_nothing_usable_remains(self, value: str) -> None:
        """Auch reservierte Geraetenamen wie ``CON`` sind unter Windows tabu."""
        assert safe_filename(value) == "Testlauf"

    def test_shortens_overly_long_names(self) -> None:
        assert len(safe_filename("A" * 300)) == 80


class TestWorkbookExport:
    """Aufbau der erzeugten Arbeitsmappe."""

    def test_file_is_named_after_the_run(self, db: Database, tmp_path: Path) -> None:
        run_id = _filled_run(db, "Vergleich WLAN")
        path = export_workbook(db, run_id, tmp_path)
        assert path.name == "Vergleich WLAN.xlsx"

    def test_second_run_with_the_same_name_does_not_overwrite(
        self, db: Database, tmp_path: Path
    ) -> None:
        """Namen sind frei waehlbar - ein Export darf keinen anderen ersetzen."""
        first = export_workbook(db, _filled_run(db, "Test"), tmp_path)
        second = export_workbook(db, _filled_run(db, "Test"), tmp_path)
        assert first != second
        assert first.exists() and second.exists()

    def test_contains_all_expected_sheets(self, db: Database, tmp_path: Path) -> None:
        path = export_workbook(db, _filled_run(db), tmp_path)
        sheets = load_workbook(path).sheetnames
        assert sheets[0] == "Uebersicht", "Die Uebersicht muss zuerst kommen."
        for expected in ("Messreihen", "Ereignisse", "Bandbreite", "Messwerte"):
            assert expected in sheets

    def test_overview_carries_run_details(self, db: Database, tmp_path: Path) -> None:
        path = export_workbook(db, _filled_run(db, "Firmware 8.02"), tmp_path)
        sheet = load_workbook(path)["Uebersicht"]
        text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row)
        assert "Firmware 8.02" in text
        assert "FRITZ!Box 7590" in text
        assert "Verfuegbarkeit" in text

    def test_no_blank_row_below_the_header(self, db: Database, tmp_path: Path) -> None:
        """Ein Filter oder eine Sortierung in Excel endet an einer leeren Zeile."""
        path = export_workbook(db, _filled_run(db, count=5), tmp_path)
        sheet = load_workbook(path)["Messwerte"]
        assert sheet.max_row == 6, "Kopfzeile plus fuenf Messwerte"
        assert any(cell.value not in (None, "") for cell in sheet[2])

    def test_local_time_column_is_added(self, db: Database, tmp_path: Path) -> None:
        """Unix-Sekunden bleiben erhalten, sind aber nicht lesbar."""
        path = export_workbook(db, _filled_run(db), tmp_path)
        sheet = load_workbook(path)["Messwerte"]
        assert [cell.value for cell in sheet[1]][-1] == "zeit_lokal"

    def test_series_sheet_lists_statistics(self, db: Database, tmp_path: Path) -> None:
        path = export_workbook(db, _filled_run(db, count=5), tmp_path)
        rows = list(load_workbook(path)["Messreihen"].iter_rows(values_only=True))
        assert rows[0] == (
            "Modul", "Messreihe", "Einheit", "Anzahl",
            "Minimum", "Mittel", "Median", "p95", "Maximum",
        )
        entry = next(row for row in rows[1:] if row[1] == "rtt_avg")
        assert entry[3] == 5
        assert entry[4] == 10.0
        assert entry[8] == 14.0

    def test_raw_sheet_is_capped_and_says_so(
        self, db: Database, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Abgeschnittene Messdaten muessen in der Datei sichtbar sein.

        Stillschweigend fehlende Werte waeren in einem Messwerkzeug der
        schlimmste Fehler: Die Datei sieht vollstaendig aus.
        """
        monkeypatch.setattr(workbook_module, "MAX_RAW_ROWS", 3)
        path = export_workbook(db, _filled_run(db, count=10), tmp_path)
        sheet = load_workbook(path)["Messwerte"]
        assert sheet.max_row == 5, "Kopfzeile, drei Messwerte, ein Hinweis"
        assert "abgeschnitten" in str(sheet.cell(row=5, column=1).value)

    def test_empty_table_gets_a_note_instead_of_an_empty_sheet(
        self, db: Database, tmp_path: Path
    ) -> None:
        path = export_workbook(db, _filled_run(db), tmp_path)
        sheet = load_workbook(path)["Ausfaelle"]
        assert "Keine Daten" in str(sheet.cell(row=1, column=1).value)

    def test_unknown_run_raises(self, db: Database, tmp_path: Path) -> None:
        with pytest.raises(ValueError, match="existiert nicht"):
            export_workbook(db, 999, tmp_path)
